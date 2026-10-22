"""Normalize checkpoints to narrow scientific tables; export/merge independently of PySCF."""
import csv
import hashlib
import json
import math
import re
from pathlib import Path

ANALYTIC={'uhf','pbe','pbe0','b3lyp'}


def record_status(record):
    status=record.get('status','missing')
    # Older checkpoints used a generic error for all downstream methods.
    if (not record.get('valid') and not record.get('scf',{}).get('valid',True)
            and (record.get('error_stage')=='scf_reference'
                 or 'No converged admissible SCF reference' in (record.get('error') or ''))):
        return 'reference_failed' if record.get('method')==record.get('source') else 'not_run_reference_failed'
    return status


def definitions(method,mode,fields):
    names=['analytic' if method in ANALYTIC else 'determinant']
    if method not in ANALYTIC and mode in ('dm','both'):
        names.append('cc_lambda' if method in ('ccsd','ccsd_t') else 'scaled_t2' if method in ('obdh','b2plyp','dh_matched') else 't2')
    return [(n,None) for n in names]+([('finite_field',f) for f in fields] if mode in ('ff','both') else [])


def scalar_rows(record,prefix=''):
    rows=[]
    for k,v in record.items():
        if isinstance(v,dict):rows.extend(scalar_rows(v,prefix+k+'.'))
        else:rows.append(dict(parameter=prefix+k,value=json.dumps(v) if isinstance(v,list) else v))
    return rows


def rows_for(meta,cases,reference):
    cfg=meta['config'];requested=meta['requested_indices'];by_index={c['index']:c for c in cases}
    results=[];energies=[];vectors=[];convergence=[];candidates=[];densities=[];fields=[];molecules=[]
    for i in requested:
        name=meta['inputs'][i]['name'];case=by_index.get(i,{})
        props=meta['inputs'][i].get('props',{})
        molecules.append(dict(molecule=name,charge=props.get('charge'),spin=props.get('spin'),
            atoms=len(props.get('geometry',[])),nao=case.get('nao'),complete=case.get('complete',False),
            threads=case.get('execution',{}).get('threads'),max_memory_MB=case.get('execution',{}).get('max_memory'),
            geometry_unit='Angstrom',point_file=f'points/{i:03d}_'+re.sub(r'[^A-Za-z0-9_.-]','_',name)+'.json'))
        zeros={r['method']:r for r in case.get('zero',[])}
        for method in cfg['methods']:
            zero=zeros.get(method,{})
            energies.append(dict(molecule=name,method=method,energy_Eh=zero.get('energy_Eh') if zero.get('valid') else None,
                last_energy_Eh=zero.get('energy_Eh'),valid=zero.get('valid',False),status=record_status(zero),
                reference=zero.get('source'),error=zero.get('error',case.get('error'))))
            for estimator,field in definitions(method,cfg['dipole'],cfg['fields']):
                datum=(next((d for d in case.get('finite_field',{}).get(method,[]) if d['field_au']==field),{})
                       if estimator=='finite_field' else zero.get('densities',{}).get(estimator,{}))
                valid=bool(datum.get('valid',False));raw=datum.get('magnitude_D')
                reason=('converged' if valid else 'step_not_converged' if datum.get('step_check')=='failed'
                        else 'invalid_field_point' if estimator=='finite_field' and datum
                        else 'lambda_not_converged' if zero.get('lambda_converged') is False
                        else 'density_invalid' if datum and datum.get('density_valid') is False
                        else 'density_error' if zero.get('density_error') and not datum
                        else record_status(zero) if not datum else 'not_converged')
                ref=reference.get(name)
                if ref is not None and (not isinstance(ref,(float,int)) or not math.isfinite(ref) or ref<0):
                    raise ValueError(f'Invalid reference dipole for {name}')
                mu=raw if valid else None
                err=mu-ref if mu is not None and ref is not None else None
                results.append(dict(molecule=name,method=method,estimator=estimator,field_au=field,
                    dipole_D=mu,reference_D=ref,error_D=err,regularized_error_pct=100*err/max(ref,1.) if err is not None else None,
                    valid=valid,status=reason))
                vector=datum.get('vector_D') or [None]*3
                vectors.append(dict(molecule=name,method=method,estimator=estimator,field_au=field,
                    mux_D=vector[0],muy_D=vector[1],muz_D=vector[2],last_magnitude_D=raw,valid=valid,
                    step_check=datum.get('step_check'),step_spread_D=datum.get('step_spread_D')))
                if estimator!='finite_field' and datum:
                    densities.append(dict(molecule=name,method=method,estimator=estimator,
                        **{k:v for k,v in datum.items() if k not in ('vector_D','magnitude_D')}))
        for r in case.get('zero',[])+case.get('field_points',[]):
            s=r.get('scf',{});iteration=r.get('iteration') or {}
            convergence.append(dict(molecule=name,method=r['method'],axis=r.get('axis') or 'zero',field_au=r.get('field_au',0),
                valid=r.get('valid',False),status=record_status(r),scf_converged=s.get('scf_converged'),
                post_converged=r.get('post_converged'),lambda_converged=r.get('lambda_converged'),
                scf_gradient=s.get('gradient'),reference_s2=s.get('s2'),ob_fia=iteration.get('effective_fia'),
                cycles=iteration.get('cycles'),delta_energy=iteration.get('delta_energy'),
                reference_overlap=s.get('reference_overlap'),ob_overlap=r.get('ob_overlap'),
                ob_stability=r.get('ob_stability'),seconds=r.get('seconds'),error=r.get('error'),density_error=r.get('density_error')))
            if r.get('axis') is not None:
                fields.append(dict(molecule=name,method=r['method'],axis=r['axis'],field_au=r['field_au'],
                    energy_Eh=r.get('energy_Eh'),valid=r.get('valid',False),scf_converged=s.get('scf_converged'),
                    post_converged=r.get('post_converged'),reference_overlap=s.get('reference_overlap'),
                    ob_overlap=r.get('ob_overlap'),error=r.get('error')))
        for r in case.get('scf_candidates',[]):
            candidates.append(dict(molecule=name,**{k:r.get(k) for k in
                ('source','guess','selected','energy_Eh','valid','scf_converged','stable','gradient','s2','density_valid','rejection_reasons','error')}))
    estimator_order=['analytic','determinant','t2','scaled_t2','cc_lambda','finite_field']
    key=lambda r:(cfg['methods'].index(r['method']),estimator_order.index(r['estimator']),
                  cfg['fields'].index(r['field_au']) if r['field_au'] is not None else -1,r['molecule'])
    results.sort(key=key)
    groups={}
    for r in results:groups.setdefault((r['method'],r['estimator'],r['field_au']),[]).append(r)
    cohorts=[{r['molecule'] for r in rs if r['regularized_error_pct'] is not None} for rs in groups.values()]
    common=set.intersection(*cohorts) if cohorts else set()
    stats=[]
    for (method,estimator,field),rs in groups.items():
        errs=[r['regularized_error_pct'] for r in rs if r['regularized_error_pct'] is not None]
        paired=[r['regularized_error_pct'] for r in rs if r['molecule'] in common]
        stats.append(dict(method=method,estimator=estimator,field_au=field,n_requested=len(rs),n_converged=sum(r['dipole_D'] is not None for r in rs),n_paired=len(errs),
            RMSE_available_pct=math.sqrt(sum(e*e for e in errs)/len(errs)) if errs else None,
            RMSE_complete_pct=math.sqrt(sum(e*e for e in errs)/len(errs)) if errs and len(errs)==len(rs) else None,
            MAE_available_pct=sum(abs(e) for e in errs)/len(errs) if errs else None,n_common=len(paired),
            RMSE_common_pct=math.sqrt(sum(e*e for e in paired)/len(paired)) if paired else None))
    return dict(Statistics=stats,Results=results,Molecules=molecules,Energies=energies,Dipole_vectors=vectors,
        Convergence=convergence,SCF_candidates=candidates,Field_points=fields,Density_checks=densities)


def column(n):
    result=''
    while n:n,r=divmod(n-1,26);result=chr(65+r)+result
    return result


def workbook_plan(tables,meta,reference_info):
    sheets=[]
    notes={
        'Statistics':'Error = 100*(mu-mu_ref)/max(mu_ref,1 D). Full RMSE stays blank until all requested pairs are valid. n_converged counts dipoles; n_paired counts reference pairs. Common cohort is across all method/estimator/field groups, saved at merge time.',
        'Results':'Accepted dipoles only. Invalid or uncomputed values stay blank. Error formulas use magnitudes; signed vectors are in Dipole_vectors.',
        'Energies':'energy_Eh is accepted; last_energy_Eh preserves failed/unconverged iterates for diagnosis.',
        'Dipole_vectors':'Signed components and raw magnitudes, including invalid iterates. No analytic value is labelled finite_field.',
        'Convergence':'Zero and field convergence. OB stability is not tested; SCF stability does not establish OB energy stationarity.',
        'SCF_candidates':'Selection uses the lowest admissible SCF energy. No spin-contamination or dipole-error gate.',
        'Field_points':'Raw E(+F)/E(-F), convergence and subspace overlaps. Finite-field formula adds nuclear dipole once.',
        'Density_checks':'AO electron counts use Tr(D S); occupations use S^(1/2) D S^(1/2). Occupation violations are diagnostic, not automatic rejection.',
        'Molecules':'Charge, spin (N_alpha - N_beta), AO size and geometry provenance. Full coordinates are in the input JSON saved by the run metadata.',
        'Config':'Initial configuration and source provenance; actual case resources are in Molecules. Geometry: Angstrom. Dipoles: Debye. Reference details recorded below.'}
    tables=dict(tables)
    tables['Config']=scalar_rows({'config':meta['config'],'provenance':meta.get('provenance',{}),'reference':reference_info,
        'definitions':{'analytic':'UHF/KS determinant expectation','t2':'D_det + d_oo(T2)+d_vv(T2), no orbital response',
        'scaled_t2':'D_det + alpha_c*delta_D_T2; OBDH is a prescription, not a proven response density',
        'cc_lambda':'CC Lambda density without SCF orbital response','finite_field':'nuclear dipole - dE_electronic/dF; actual field calculations'}})
    # Put the concise main results first, then aggregate statistics and diagnostics.
    order=['Results','Statistics','Energies','Dipole_vectors','Convergence','SCF_candidates','Field_points','Density_checks','Molecules','Config']
    for name in order:
        rows=tables[name]
        if not rows and name not in ('Results','Statistics','Config'):continue
        headers=list(dict.fromkeys(k for r in rows for k in r)) or ['status']
        matrix=[[r.get(h) for h in headers] for r in rows]
        forms={}
        if name=='Results':
            for j in range(len(rows)):
                n=j+5
                forms[f'G{n}']=f'=IF(AND(ISNUMBER(E{n}),ISNUMBER(F{n})),E{n}-F{n},"")'
                forms[f'H{n}']=f'=IF(ISNUMBER(G{n}),100*G{n}/MAX(F{n},1),"")'
        if name=='Statistics':
            data=tables['Results']
            for j,r in enumerate(rows):
                indices=[n+5 for n,v in enumerate(data) if (v['method'],v['estimator'],v['field_au'])==(r['method'],r['estimator'],r['field_au'])]
                if not indices:continue
                first,last=min(indices),max(indices);rng=f"'Results'!H{first}:H{last}";n=j+5
                mu_range=f"'Results'!E{first}:E{last}"
                forms[f'E{n}']=f'=COUNT({mu_range})'
                forms[f'F{n}']=f'=COUNT({rng})'
                forms[f'G{n}']=f'=IF(F{n}=0,"",SQRT(SUMSQ({rng})/F{n}))'
                forms[f'H{n}']=f'=IF(AND(F{n}=D{n},F{n}>0),G{n},"")'
                forms[f'I{n}']=f'=IF(F{n}=0,"",(SUMIF({rng},">=0",{rng})-SUMIF({rng},"<0",{rng}))/F{n})'
        sheets.append(dict(name=name,note=notes[name],headers=headers,rows=matrix,formulas=forms))
    return dict(title='Dipole 152 benchmark',sheets=sheets)


def export_hpc_xlsx(plan,path):
    """Cluster-side adapter using the existing SIE Excel dependency (openpyxl).

    Workbook layout/formulas come exclusively from the neutral workbook plan,
    also consumable by artifact-tool for review outside the cluster.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();wb.remove(wb.active)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    for spec in plan['sheets']:
        ws=wb.create_sheet(spec['name']);ws.sheet_view.showGridLines=False
        ncols=len(spec['headers']);last=max(ncols,2)
        ws.cell(1,1,plan['title']+' — '+spec['name']);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last)
        ws.cell(2,1,spec['note']);ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=last)
        ws.row_dimensions[1].height=26;ws.row_dimensions[2].height=42;ws.row_dimensions[4].height=34
        for i,h in enumerate(spec['headers'],1):
            c=ws.cell(4,i,h);c.font=Font(name='Arial',size=11,bold=True)
            c.alignment=Alignment(wrap_text=True,vertical='center')
            c.border=Border(bottom=Side(style='thin',color='777777'))
            ws.column_dimensions[get_column_letter(i)].width=24 if h.endswith('_Eh') else 38 if h in ('error','value') else 24 if h in ('status','estimator','source','reference','parameter','ob_stability') else 18
        for j,row in enumerate(spec['rows'],5):
            for i,value in enumerate(row,1):
                if isinstance(value,(dict,list)):value=json.dumps(value)
                c=ws.cell(j,i,value);c.font=Font(name='Arial',size=11)
                c.alignment=Alignment(vertical='top',wrap_text=isinstance(value,str))
                header=spec['headers'][i-1]
                if isinstance(value,(float,int)) and not isinstance(value,bool):
                    c.number_format='General' if header=='value' else '0' if isinstance(value,int) else '0.000000000000' if header.endswith('_Eh') else '0.000000' if header.endswith('_D') else '0.000' if 'pct' in header else '0.00E+00' if any(k in header for k in ('gradient','fia','field_au','delta_energy','hermiticity')) else '0.0000'
        for j,row in enumerate(spec['rows'],5):
            lines=max((math.ceil(len(str(v))/max(10,ws.column_dimensions[get_column_letter(i)].width-2)) for i,v in enumerate(row,1) if isinstance(v,str)),default=1)
            ws.row_dimensions[j].height=max(18,15*min(lines,12))
        for address,formula in spec['formulas'].items():
            ws[address]=formula
            h=spec['headers'][ws[address].column-1]
            ws[address].number_format='0' if h.startswith('n_') else '0.000' if 'pct' in h else '0.000000'
        ws['A1'].font=Font(name='Arial',size=14,bold=True);ws['A2'].alignment=Alignment(wrap_text=True,vertical='center')
        ws.freeze_panes='C5';ws.auto_filter.ref=f'A4:{get_column_letter(ncols)}{4+len(spec["rows"])}'
        ws.print_title_rows='1:4';ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_name(path.stem+'.tmp.xlsx');wb.save(temp);temp.replace(path)


def collect(outdir,ref_json=None,output_xlsx=None,excel=True):
    outdir=Path(outdir);meta=json.loads((outdir/'dipole.meta.json').read_text())
    cases=[]
    for path in sorted((outdir/'points').glob('*.json')):
        case=json.loads(path.read_text())
        if case.get('fingerprint')!=meta['fingerprint']:raise RuntimeError(f'Mixed configuration: {path}')
        cases.append(case)
    reference_info={};reference={}
    ref_json=ref_json or meta.get('ref_json')
    if ref_json:
        path=Path(ref_json);raw=json.loads(path.read_text());reference=raw.get('values',raw)
        reference_info=dict(path=str(path.resolve()),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),metadata=raw.get('metadata',{}))
    tables=rows_for(meta,cases,reference)
    for name,rows in tables.items():
        headers=list(dict.fromkeys(k for r in rows for k in r)) or ['status']
        temp=outdir/(name+'.csv.tmp')
        with temp.open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=headers);w.writeheader()
            w.writerows({k:json.dumps(v) if isinstance(v,(dict,list)) else v for k,v in r.items()} for r in rows)
        temp.replace(outdir/(name+'.csv'))
    plan=workbook_plan(tables,meta,reference_info)
    (outdir/'dipole.workbook.json').write_text(json.dumps(plan,indent=2,allow_nan=False)+'\n')
    if excel:export_hpc_xlsx(plan,output_xlsx or outdir/'dipole.xlsx')
    complete=sum(c.get('complete',False) for c in cases if c['index'] in meta['requested_indices'])
    text=f'Requested molecules: {len(meta["requested_indices"])}\nComplete: {complete}\n'
    for row in tables['Statistics']:text+=json.dumps(row)+'\n'
    (outdir/'summary.txt').write_text(text)
    print(text,flush=True)
    return tables
