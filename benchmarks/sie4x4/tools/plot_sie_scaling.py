#!/usr/bin/env python3
"""Plot SIE4x4 OBDH scaling from the vertical Excel summary template.

Dependencies: numpy, matplotlib, openpyxl
Example (direct invocation from any working directory):
    python /home/giahuy/Code/job/benchmarks/sie4x4/tools/plot_sie_scaling.py \
        --input "/data/giahuy/Result/OBDH/SIE_DFT/SIE_scaling_template.xlsx"

By default, SIE_UHF_scaling.png is saved beside the input Excel file.
Use --output /absolute/path/figure.pdf to choose another destination.

Reads only the Excel file, never the scaling directories. Uses Nang_luong
column E (E_total, Ha), Bo_scaling columns A:E (ID, Alpha, method, basis,
init_source), M4 (Ha-to-kcal), and L:N (system, R, reference De).
Merged IDs and system names are propagated within their blocks.
Only obdh with init_source=uhf is plotted. Only enter valid, converged
selected energies in this template: a numeric energy alone cannot prove
convergence. Blank/text energies remain missing, never zero. Overall MUE
requires all 16 errors; system MUE requires all four distances.
Calculations use raw input energies, not cached Excel formula results.
The workbook is never modified. Existing copying errors remain input errors.
"""

import argparse
import re
from numbers import Real

from openpyxl import load_workbook
from pathlib import Path

import numpy as np
import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

SYSTEMS = ('H2_plus_He', 'He2_plus', 'NH3_2_plus', 'H2O_2_plus')
NAMES = ('H₂⁺···He', 'He₂⁺', '(NH₃)₂⁺', '(H₂O)₂⁺')
ALPHAS = np.arange(1, 10) / 10


def numeric(value):
    return isinstance(value, Real) and not isinstance(value, bool) and np.isfinite(value)


def read_scan(path):
    # Formula evaluation is deliberately unnecessary: input cells must be raw values.
    workbook = load_workbook(path, data_only=False)
    required = {'Bo_scaling', 'Nang_luong'}
    if not required.issubset(workbook.sheetnames):
        raise ValueError('Cần các sheet Bo_scaling và Nang_luong theo mẫu tổng hợp dạng dọc.')
    meta = workbook['Bo_scaling']
    energy_sheet = workbook['Nang_luong']
    conversion = meta['M4'].value
    if not numeric(conversion) or conversion <= 0:
        raise ValueError('Bo_scaling!M4 phải chứa hệ số Ha → kcal/mol dạng số dương.')
    references = {}
    system = None
    for row in meta.iter_rows(min_row=8, min_col=12, max_col=14, values_only=True):
        name, distance, value = row
        if name is not None:
            system = str(name).strip()
        if system in SYSTEMS and numeric(distance) and numeric(value):
            key = (system, float(distance))
            if key in references:
                raise ValueError(f'Reference bị trùng: {key}')
            references[key] = float(value)
    distances = (1.0, 1.25, 1.5, 1.75)
    for system in SYSTEMS:
        for distance in distances:
            if (system, distance) not in references:
                raise ValueError(f'Thiếu reference dạng số cho {system}, R={distance} tại Bo_scaling!L:N.')
    records = {}
    for row in meta.iter_rows(min_row=8, max_col=5, values_only=True):
        identifier, alpha, method, basis, source = row
        if not identifier or str(method).strip().lower() != 'obdh' or str(source).strip().lower() != 'uhf':
            continue
        match = re.fullmatch(r'\s*\(\s*([\d.]+)\s*,\s*([\d.]+)\s*\)\s*', str(alpha))
        if not match:
            raise ValueError(f'Alpha không hợp lệ cho {identifier}: {alpha!r}')
        hf, pt2 = map(float, match.groups())
        if not all(any(np.isclose(v, a) for a in ALPHAS) for v in (hf, pt2)):
            raise ValueError(f'{identifier}: Alpha phải thuộc lưới 0.1–0.9.')
        key = (round(hf * 10) - 1, round(pt2 * 10) - 1)
        if identifier in records or any(r['key'] == key for r in records.values()):
            raise ValueError(f'ID hoặc Alpha bị trùng trong các bộ OBDH/UHF: {identifier}, {alpha}')
        records[identifier] = {'key': key, 'basis': basis, 'energies': {}, 'seen': set()}
    if not records:
        raise ValueError('Không có bộ method=obdh, init_source=uhf trong Bo_scaling!C:E.')
    identifier = system = None
    for row_number, row in enumerate(energy_sheet.iter_rows(min_row=8, max_col=5, values_only=True), start=8):
        record_id, _, name, distance, value = row
        if record_id is not None:
            identifier = record_id
            system = None
        if name is not None:
            system = str(name).strip()
        if identifier not in records or system not in SYSTEMS:
            continue
        if numeric(distance):
            point = float(distance)
            if point not in distances:
                continue
        elif str(distance).strip().lower() in {'r_dl', 'r dl', 'dl', 'dissociation_limit', 'r∞', 'r_∞'}:
            point = 'DL'
        else:
            continue
        record = records[identifier]
        key = (system, point)
        if key in record['seen']:
            raise ValueError(f'Nang_luong dòng {row_number}: trùng {identifier}, {key}.')
        record['seen'].add(key)
        if isinstance(value, str) and value.startswith('='):
            raise ValueError(f'Nang_luong!E{row_number}: cần E_total dạng số, không phải công thức.')
        if numeric(value):
            record['energies'][key] = float(value)
    mue = np.full((9, 9), np.nan)
    system_mue = np.full((9, 9, 4), np.nan)
    counts = np.zeros((9, 9), dtype=int)
    configs = {}
    for record in records.values():
        i, j = record['key']
        energies = record['energies']
        errors = np.full((4, 4), np.nan)
        for s, system in enumerate(SYSTEMS):
            limit = energies.get((system, 'DL'))
            if limit is None:
                continue
            for p, distance in enumerate(distances):
                energy = energies.get((system, distance))
                if energy is not None:
                    de = (limit - energy) * conversion
                    errors[s, p] = de - references[system, distance]
            if np.isfinite(errors[s]).all():
                system_mue[i, j, s] = np.abs(errors[s]).mean()
        counts[i, j] = np.isfinite(errors).sum()
        if counts[i, j] == 16:
            mue[i, j] = np.abs(errors).mean()
        if record['basis']:
            configs[i, j] = {'basis': record['basis']}
    workbook.close()
    return mue, system_mue, counts, configs, len(records)


def draw(mue, system_mue, counts, configs, hf, output, vmax, dpi):
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10,
                         'axes.spines.top': False, 'axes.spines.right': False})
    fig, (ax, bx) = plt.subplots(1, 2, figsize=(15, 6.8),
        gridspec_kw={'width_ratios': [1.08, 1]}, layout='constrained')
    cmap = plt.colormaps['YlOrRd'].copy()
    cmap.set_bad('#E5E7EB')
    im = ax.imshow(np.ma.masked_invalid(mue), origin='lower', cmap=cmap,
                   vmin=0, vmax=vmax, aspect='equal')
    for i in range(9):
        for j in range(9):
            label = f'{mue[i, j]:.2f}' if np.isfinite(mue[i, j]) else f'N={counts[i, j]}'
            ax.text(j, i, label, ha='center', va='center', fontsize=8.7,
                    color='white' if mue[i, j] > 0.73 * vmax else '#111827')
    if np.isfinite(mue).any():
        bi, bj = np.unravel_index(np.nanargmin(mue), mue.shape)
        ax.add_patch(Rectangle((bj - .5, bi - .5), 1, 1, fill=False,
                              lw=2.4, edgecolor='#111827'))
    labels = [f'{v:.1f}' for v in ALPHAS]
    ax.set_xticks(range(9), labels)
    ax.set_yticks(range(9), labels)
    ax.set_xlabel('alpha_c')
    ax.set_ylabel('alpha_x')
    ax.set_title('A. MUE \n Grey points: discovegence->missing results',
                 fontweight='bold', pad=12)
    fig.colorbar(im, ax=ax, shrink=.78, label='MUE (kcal/mol)')
    index = round(hf * 10) - 1
    colors = ('#0072B2', '#E69F00', '#009E73', '#CC79A7')
    for s, (name, color) in enumerate(zip(NAMES, colors)):
        bx.plot(ALPHAS, system_mue[index, :, s], marker='o', lw=1.8,
                markersize=4, color=color, label=name)
    bx.plot(ALPHAS, mue[index], color='#111827', ls='--', lw=2,
            marker='s', markersize=4, label='Results')
    if np.isfinite(mue[index]).any():
        j = np.nanargmin(mue[index])
        bx.axvline(ALPHAS[j], color='#6B7280', lw=1, ls=':')
        bx.scatter([ALPHAS[j]], [mue[index, j]], s=110, facecolors='none',
                   edgecolors='#111827', linewidths=1.5, zorder=5)
    bx.set_xticks(ALPHAS)
    bx.set_ylim(bottom=0)
    bx.set_xlabel(f'OBMP2 scaling (fixed HF scaling = {hf:.1f})')
    bx.set_ylabel('MUE (kcal/mol)')
    bx.grid(axis='y', alpha=.2)
    # bx.set_title(f'B. Phân rã MUE khi hệ số HF = {hf:.1f}\nCân bằng sai số giữa bốn hệ',
    #              fontweight='bold', pad=12)
    bx.legend(frameon=False, fontsize=9, loc='upper right')
    bases = {str(c['basis']) for c in configs.values() if 'basis' in c}
    grids = {str(c['grid']) for c in configs.values() if 'grid' in c}
    details = ''
    if len(bases) == 1:
        details += ' · ' + next(iter(bases))
    if len(grids) == 1:
        details += ' · grid ' + next(iter(grids))
    # fig.suptitle('SIE4x4 · OBDH từ UHF-init' + details, fontsize=15, fontweight='bold')
    # footer = 'Nguồn: file Excel tổng hợp · Ô trống/chữ: thiếu dữ liệu · N: số điểm hợp lệ/16.'
    # footer += '\nMUE tính lại từ E_total và reference trong Excel; khung đen: MUE thấp nhất.'
    # fig.supxlabel(footer, fontsize=8.5)
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, dpi=dpi, facecolor='white')
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input', type=Path, required=True, help='File Excel tổng hợp .xlsx.')
    parser.add_argument('--output', type=Path, default=None,
                        help='Định dạng PNG hoặc PDF tùy phần mở rộng.')
    parser.add_argument('--hf', type=float, default=.9, help='Hệ số HF của biểu đồ B; 0.1–0.9.')
    parser.add_argument('--vmax', type=float, default=30, help='Giới hạn trên thang màu MUE.')
    parser.add_argument('--dpi', type=int, default=180)
    args = parser.parse_args()
    if not any(np.isclose(args.hf, value) for value in ALPHAS):
        parser.error('--hf phải là một trong 0.1, 0.2, ..., 0.9')
    if args.vmax <= 0 or args.dpi <= 0:
        parser.error('--vmax và --dpi phải dương')
    try:
        mue, system_mue, counts, configs, found = read_scan(args.input.expanduser())
        output = args.output.expanduser() if args.output else args.input.expanduser().with_name('SIE_UHF_scaling.png')
        draw(mue, system_mue, counts, configs, args.hf, output, args.vmax, args.dpi)
    except (OSError, ValueError) as exc:
        parser.exit(1, f'Lỗi: {exc}\n')
    print(f'Đã đọc {found} bộ OBDH/UHF từ Excel; {np.isfinite(mue).sum()} bộ đủ 16 điểm.')
    if np.isfinite(mue).any():
        i, j = np.unravel_index(np.nanargmin(mue), mue.shape)
        print(f'MUE thấp nhất: Alpha=({ALPHAS[i]:.1f},{ALPHAS[j]:.1f}), {mue[i, j]:.6f} kcal/mol')
    print(f'Đã lưu: {output.resolve()}')


if __name__ == '__main__':
    main()
